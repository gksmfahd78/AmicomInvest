from pathlib import Path
import json,zipfile,hashlib,unicodedata
import pymupdf as fitz
from pptx import Presentation
ROOT=Path(__file__).resolve().parents[2];out=ROOT/'presentations/output';review=ROOT/'presentations/review'
report=json.loads((review/'render-review.json').read_text());assert not report['missingText'] and not report['outsideText']
checks=[]
norm=lambda s:''.join(unicodedata.normalize('NFKC',s).split())
for path in sorted(out.glob('*.pptx')):
 p=Presentation(path);pdf=fitz.open(path.with_suffix('.pdf'));assert len(p.slides)==len(pdf)
 for idx,(slide,page) in enumerate(zip(p.slides,pdf)):
  rendered=norm(page.get_text())
  for shape in slide.shapes:
   if shape.has_text_frame:assert norm(shape.text) in rendered,(path.name,idx+1,shape.text)
  assert len(slide.notes_slide.notes_text_frame.text)>100
  assert sum(sh.has_text_frame for sh in slide.shapes)>3
 checks.append({'file':path.name,'slides':len(p.slides),'pdfMatched':True,'notesPresent':True,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
allpdf=fitz.open(out/'AMICOM_Investment_Course.pdf');preview=fitz.open()
for i in [0,25,26,38,49,69,72,85]:preview.insert_pdf(allpdf,from_page=i,to_page=i)
preview.save(out/'AMICOM_Design_Preview.pdf',garbage=4,deflate=True)
# Representative slide image for quick inspection; same rendered artifact as the delivered PPT.
allpdf[25].get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False).save(out/'AMICOM_Slide_Preview.png')
archive=ROOT/'presentations/AMICOM_8Week_PPT_Package.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
 for path in sorted(out.rglob('*')):
  if path.is_file() and path.name!='manifest.json':z.write(path,'AMICOM_PPT/'+str(path.relative_to(out)))
with zipfile.ZipFile(archive) as z:assert z.testzip() is None
result={'presentations':checks,'renderedAndVerified':True,'designReview':'All 90 slides reviewed in rendered contact sheets; key charts reviewed at full resolution.','package':str(archive),'packageBytes':archive.stat().st_size,'fontPack':'Noto Sans CJK Regular/Bold with redistribution license','nativePowerPointTested':False,'renderer':'LibreOffice Impress'}
(review/'delivery-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps({'pptxFiles':len(checks),'allTextPresent':True,'allNotesPresent':True,'zipMB':round(archive.stat().st_size/1024/1024,1),'archive':str(archive)},ensure_ascii=False))
